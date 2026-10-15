import io
import unittest
from datetime import date

from docx import Document
from openpyxl import load_workbook
from lunch_reports import (parse_lunch_events, lunch_tables, export_lunch_docx,
                           export_lunch_pdf, export_lunch_xlsx)

DAY = date(2026, 10, 14)
NAME = 'Lunch Test Person'


def scans(times, employee_id='101', name=NAME, day=DAY):
    return [{'employeeNoString': employee_id, 'name': name,
             'time': f'{day.isoformat()}T{time}:00+03:00'} for time in times]


def site():
    return {'dept_order': ['Shop'], 'employees': {'Shop': {
        'label': 'SHOP', 'members': [NAME, 'Absent Person'],
        'member_ids': {NAME: '101', 'Absent Person': '102'}}}}


class LunchClockingTests(unittest.TestCase):
    def test_requested_four_clockings_produce_seven_net_hours(self):
        row = parse_lunch_events(scans(['08:00', '13:00', '14:00', '16:00']), DAY, DAY)[0]
        self.assertEqual([row[key] for key in ['check_in', 'lunch_out', 'lunch_in', 'check_out']],
                         ['08:00', '13:00', '14:00', '16:00'])
        self.assertEqual(row['lunch_minutes'], 60)
        self.assertEqual(row['net_minutes'], 420)
        self.assertEqual(row['net_hours'], '7h 00m')
        self.assertTrue(row['complete'])

    def test_missing_return_does_not_turn_departure_into_lunch_return(self):
        row = parse_lunch_events(scans(['08:00', '13:00', '16:00']), DAY, DAY)[0]
        self.assertEqual(row['lunch_in'], '')
        self.assertEqual(row['check_out'], '16:00')
        self.assertEqual(row['missing'], 'Lunch In')
        self.assertIsNone(row['net_minutes'])
        self.assertFalse(row['complete'])

    def test_only_arrival_never_fabricates_lunch_or_departure(self):
        row = parse_lunch_events(scans(['08:00', '08:01', '10:00']), DAY, DAY)[0]
        self.assertEqual(row['check_in'], '08:00')
        self.assertEqual((row['lunch_out'], row['lunch_in'], row['check_out']), ('', '', ''))
        self.assertEqual(row['net_hours'], '')

    def test_lunch_scans_without_arrival_are_flagged(self):
        row = parse_lunch_events(scans(['13:05', '14:00', '16:02']), DAY, DAY)[0]
        self.assertEqual(row['check_in'], '')
        self.assertEqual(row['lunch_out'], '13:05')
        self.assertEqual(row['missing'], 'Check In')

    def test_duplicates_and_unsorted_scans_do_not_change_slots(self):
        rows = parse_lunch_events(scans(['16:10', '14:20', '13:10', '08:03', '08:03', '13:11', '14:21']), DAY, DAY)
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row['net_hours'], '6h 57m')
        self.assertEqual(row['lunch_minutes'], 70)

    def test_local_dates_and_date_limits(self):
        events = scans(['08:00'], day=date(2026, 10, 13))
        events.append({'employeeNo': '101', 'name': NAME, 'time': '2026-10-14T05:00:00Z'})
        rows = parse_lunch_events(events, DAY, DAY)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['check_in'], '08:00')
        self.assertEqual(rows[0]['date'], DAY.isoformat())

    def test_seconds_match_the_displayed_minute_precision(self):
        events = scans(['08:00', '13:00', '14:00', '16:00'])
        events[0]['time'] = '2026-10-14T08:00:59+03:00'
        self.assertEqual(parse_lunch_events(events, DAY, DAY)[0]['net_minutes'], 420)

    def test_malformed_events_do_not_create_records(self):
        self.assertEqual(parse_lunch_events([{'name': NAME, 'time': 'bad'},
                                            {'employeeNo': '101', 'time': '2026-10-14T08:00:00'}], DAY, DAY), [])

    def test_absent_staff_and_summary_of_complete_and_partial_days(self):
        next_day = date(2026, 10, 15)
        events = scans(['08:00', '13:00', '14:00', '16:00']) + scans(['08:30', '16:00'], day=next_day)
        rows = parse_lunch_events(events, DAY, next_day)
        daily, summary = lunch_tables(rows, site(), DAY, next_day)
        self.assertEqual(len(daily), 2)
        self.assertEqual(daily[0][2][1]['Status'], 'Absent')
        self.assertEqual(daily[1][2][0]['Status'], 'Incomplete')
        record = summary[0][2][0]
        self.assertEqual((record['Present'], record['Absent'], record['Complete'], record['Incomplete']), (2, 0, 1, 1))
        self.assertEqual(record['Net Hours'], '7h 00m')
        self.assertEqual(record['Avg In'], '08:15')
        self.assertEqual(summary[0][2][1]['Absent'], 2)
        self.assertEqual(summary[0][2][1]['Net Hours'], '')

    def test_deleted_or_unlisted_event_is_not_added_to_roster_tables(self):
        rows = parse_lunch_events(scans(['08:00'], employee_id='999', name='Unlisted Person'), DAY, DAY)
        daily, summary = lunch_tables(rows, site(), DAY, DAY)
        self.assertNotIn('Unlisted Person', str(daily) + str(summary))


class LunchExportTests(unittest.TestCase):
    def test_daily_and_summary_exports_contain_real_lunch_values(self):
        rows = parse_lunch_events(scans(['08:00', '13:00', '14:00', '16:00']), DAY, DAY)
        daily, summary = lunch_tables(rows, site(), DAY, DAY)
        for tables in [daily, summary]:
            word = Document(io.BytesIO(export_lunch_docx('BUGURUNI LUNCH REPORT', '14 October 2026', tables)))
            text = ' '.join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
            self.assertIn(NAME, text)
            self.assertIn('13:00', text)
            self.assertIn('14:00', text)
            self.assertIn('7h 00m', text)
            excel = load_workbook(io.BytesIO(export_lunch_xlsx('BUGURUNI LUNCH REPORT', '14 October 2026', tables)))
            cells = [cell.value for row in excel.active for cell in row]
            for value in [NAME, '13:00', '14:00', '7h 00m']:
                self.assertIn(value, cells)
            pdf = export_lunch_pdf('BUGURUNI LUNCH REPORT', '14 October 2026', tables)
            self.assertTrue(pdf.startswith(b'%PDF'))

    def test_terminal_text_is_not_an_excel_formula(self):
        tables = [('SHOP', ['Employee Name', 'ID'], [{'Employee Name': '=1+1', 'ID': '=2+2'}])]
        workbook = load_workbook(io.BytesIO(export_lunch_xlsx('Title', 'Subtitle', tables)))
        self.assertEqual(workbook.active['A8'].value, '=1+1')
        self.assertEqual(workbook.active['A8'].data_type, 's')


if __name__ == '__main__':
    unittest.main()
