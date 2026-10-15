"""Optional Buguruni day report: arrival, lunch out/back, and departure."""
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from io import BytesIO
from xml.sax.saxutils import escape

from roster_store import normalized

LOCAL_TIMEZONE = timezone(timedelta(hours=3))
SCHEDULE = 'Day shift: 08:00-16:00 | Lunch: 13:00-14:00'
SLOTS = [('check_in', 'Check In'), ('lunch_out', 'Lunch Out'),
         ('lunch_in', 'Lunch In'), ('check_out', 'Check Out')]
DAILY_COLUMNS = ['Status', 'Employee Name', 'ID', 'Check In', 'Lunch Out',
                 'Lunch In', 'Check Out', 'Lunch (min)', 'Net Hours', 'Missing Clockings']
SUMMARY_COLUMNS = ['Employee Name', 'ID', 'Present', 'Absent', 'Complete', 'Incomplete',
                   'Avg In', 'Avg Lunch Out', 'Avg Lunch In', 'Avg Out', 'Net Hours']


def duration_text(minutes):
    hours, minutes = divmod(minutes, 60)
    return f'{hours}h {minutes:02d}m'


def parse_lunch_events(events, start_date, end_date):
    """Classify scans by time, so a missing scan never shifts the other slots.

    The first scan before 13:00 is arrival; the first in [13:00,14:00)
    is lunch out; the first in [14:00,16:00) is lunch back; the first at
    or after 16:00 is departure. Times are actual scans, never defaults.
    This optional report uses calendar days, not overnight shift allocation.
    """
    grouped = defaultdict(list)
    for event in events:
        employee_id = str(event.get('employeeNoString') or event.get('employeeNo') or '').strip()
        name = str(event.get('name') or '').strip()
        if not employee_id or not name:
            continue
        try:
            timestamp = datetime.fromisoformat(str(event.get('time', '')).replace('Z', '+00:00'))
        except ValueError:
            continue
        timestamp = timestamp.replace(tzinfo=LOCAL_TIMEZONE) if timestamp.tzinfo is None else timestamp.astimezone(LOCAL_TIMEZONE)
        timestamp = timestamp.replace(second=0, microsecond=0)
        if start_date <= timestamp.date() <= end_date:
            grouped[(employee_id, normalized(name), timestamp.date())].append((timestamp, name))

    rows = []
    for (employee_id, _, day), scans in grouped.items():
        scans.sort()
        selected = {}
        for timestamp, name in scans:
            minute = timestamp.hour * 60 + timestamp.minute
            slot = ('check_in' if minute < 13 * 60 else 'lunch_out' if minute < 14 * 60
                    else 'lunch_in' if minute < 16 * 60 else 'check_out')
            selected.setdefault(slot, timestamp)
        missing = [label for key, label in SLOTS if key not in selected]
        complete = not missing
        lunch_minutes = None
        if 'lunch_out' in selected and 'lunch_in' in selected:
            lunch_minutes = int((selected['lunch_in'] - selected['lunch_out']).total_seconds() // 60)
        net_minutes = None
        if complete:
            net_minutes = int(((selected['lunch_out'] - selected['check_in'])
                               + (selected['check_out'] - selected['lunch_in'])).total_seconds() // 60)
        rows.append({'employee_id': employee_id, 'name': scans[0][1], 'date': day.isoformat(),
                     **{key: selected[key].strftime('%H:%M') if key in selected else '' for key, _ in SLOTS},
                     'complete': complete, 'missing': ', '.join(missing),
                     'lunch_minutes': lunch_minutes, 'net_minutes': net_minutes,
                     'net_hours': duration_text(net_minutes) if net_minutes is not None else ''})
    return sorted(rows, key=lambda row: (row['date'], normalized(row['name']), row['employee_id']))


def average_time(rows, key):
    minutes = [int(row[key][:2]) * 60 + int(row[key][3:5]) for row in rows if row.get(key)]
    if not minutes:
        return ''
    mean = sum(minutes) // len(minutes)
    return f'{mean // 60:02d}:{mean % 60:02d}'


def lunch_tables(rows, site, start_date, end_date):
    """Use the verified current roster for both absent and scanned staff."""
    by_employee = defaultdict(list)
    for row in rows:
        by_employee[(row['employee_id'], normalized(row['name']))].append(row)
    order = list(site['dept_order'])
    # The home manager is shown only when he has a scan, as in normal reports.
    for key, department in site['employees'].items():
        if key not in order and not department.get('visitor_only') and any(
                normalized(member) == normalized(row['name'])
                for member in department.get('members', []) for row in rows):
            order.append(key)
    total_days = (end_date - start_date).days + 1
    daily, summary = [], []
    members_by_dept = {}
    for key in order:
        department = site['employees'][key]
        if department.get('visitor_only'):
            continue
        members = []
        for name in department.get('members', []):
            employee_id = str(department.get('member_ids', {}).get(name, ''))
            records = by_employee.get((employee_id, normalized(name)), [])
            members.append((name, employee_id, records))
        if members:
            members_by_dept[key] = members
            summary_rows = []
            for name, employee_id, records in members:
                complete = sum(record['complete'] for record in records)
                total = sum(record['net_minutes'] for record in records if record['net_minutes'] is not None)
                summary_rows.append({'Employee Name': name, 'ID': employee_id,
                                     'Present': len(records), 'Absent': total_days - len(records),
                                     'Complete': complete, 'Incomplete': len(records) - complete,
                                     'Avg In': average_time(records, 'check_in'),
                                     'Avg Lunch Out': average_time(records, 'lunch_out'),
                                     'Avg Lunch In': average_time(records, 'lunch_in'),
                                     'Avg Out': average_time(records, 'check_out'),
                                     'Net Hours': duration_text(total) if complete else ''})
            summary.append((department['label'], SUMMARY_COLUMNS, summary_rows))
    day = start_date
    while day <= end_date:
        for key, members in members_by_dept.items():
            daily_rows = []
            for name, employee_id, records in members:
                record = next((row for row in records if row['date'] == day.isoformat()), None)
                daily_rows.append({'Status': 'Complete' if record and record['complete'] else 'Incomplete' if record else 'Absent',
                                   'Employee Name': name, 'ID': employee_id,
                                   **{label: record[key] if record else '' for key, label in SLOTS},
                                   'Lunch (min)': record['lunch_minutes'] if record and record['lunch_minutes'] is not None else '',
                                   'Net Hours': record['net_hours'] if record else '',
                                   'Missing Clockings': record['missing'] if record else ''})
            daily.append((f"{day.strftime('%A, %d %B %Y')} — {site['employees'][key]['label']}", DAILY_COLUMNS, daily_rows))
        day += timedelta(days=1)
    return daily, summary


def column_weights(columns):
    return [4 if name == 'Employee Name' else 3.2 if name == 'Missing Clockings'
            else 1.6 if name == 'Incomplete' else 1.45 if name == 'Status' else 1.2 for name in columns]


def export_lunch_docx(title, subtitle, tables, logo_path=None):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.left_margin = section.right_margin = Cm(1.2)
    section.top_margin = section.bottom_margin = Cm(1.2)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(9)
    normal.paragraph_format.space_after = Pt(2)
    for name in ['Title', 'Heading 1', 'Heading 2']:
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(20 if name == 'Title' else 11)
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    if logo_path:
        document.add_picture(logo_path, width=Cm(3))
    document.add_heading(title, 0)
    document.add_paragraph(subtitle)
    document.add_paragraph(SCHEDULE)
    document.add_paragraph('Blank clockings mean no scan was recorded. Net hours require all four clockings.')
    for heading, columns, rows in tables:
        document.add_heading(heading, 2)
        table = document.add_table(rows=1, cols=len(columns))
        table.style = 'Table Grid'
        table.autofit = False
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            line = OxmlElement('w:' + edge)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                line.set(qn('w:' + key), value)
            borders.append(line)
        table._tbl.tblPr.append(borders)
        weights = column_weights(columns)
        widths = [27.3 * weight / sum(weights) for weight in weights]
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for cell, label in zip(table.rows[0].cells, columns):
            cell.text = label
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '8B0000')
            cell._tc.get_or_add_tcPr().append(shading)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string('FFF3CD')
        for record in rows:
            cells = table.add_row().cells
            for cell, label in zip(cells, columns):
                cell.text = str(record.get(label, ''))
            background = 'EEEEEE' if record.get('Status') == 'Absent' else 'FFF3CD' if record.get('Status') == 'Incomplete' else None
            if background:
                for cell in cells:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), background)
                    cell._tc.get_or_add_tcPr().append(shading)
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell, label, width in zip(row.cells, columns, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                margins = OxmlElement('w:tcMar')
                for edge in ['top', 'bottom', 'left', 'right']:
                    margin = OxmlElement('w:' + edge)
                    margin.set(qn('w:w'), '60')
                    margin.set(qn('w:type'), 'dxa')
                    margins.append(margin)
                cell._tc.get_or_add_tcPr().append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if label in ['Employee Name', 'Missing Clockings'] else WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph()
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_lunch_pdf(title, subtitle, tables, logo_path=None):
    from pathlib import Path
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=12 * mm,
                                rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    font_directory = Path(reportlab.__file__).parent / 'fonts'
    for name, filename in [('LunchSans', 'Vera.ttf'), ('LunchSansBold', 'VeraBd.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_directory / filename)))
    body = ParagraphStyle('Body', fontName='LunchSans', fontSize=8, leading=11)
    heading_style = ParagraphStyle('Heading', parent=body, fontName='LunchSansBold', fontSize=11, leading=14, textColor=colors.HexColor('#8B0000'))
    header = ParagraphStyle('Header', parent=body, fontName='LunchSansBold', textColor=colors.HexColor('#FFF3CD'))
    story = []
    if logo_path:
        logo = Image(logo_path)
        ratio = 75 / logo.imageWidth
        logo.drawWidth, logo.drawHeight = 75, logo.imageHeight * ratio
        story.append(logo)
    for text in [title, subtitle, SCHEDULE, 'Blank clockings mean no scan was recorded. Net hours require all four clockings.']:
        story.extend([Paragraph(escape(text), heading_style if text == title else body), Spacer(1, 5)])
    for heading, columns, rows in tables:
        story.extend([Spacer(1, 8), Paragraph(escape(heading), heading_style), Spacer(1, 4)])
        data = [[Paragraph(escape(label), header) for label in columns]]
        data.extend([[Paragraph(escape(str(record.get(label, ''))), body) for label in columns] for record in rows])
        weights = column_weights(columns)
        table = Table(data, colWidths=[document.width * weight / sum(weights) for weight in weights], repeatRows=1)
        style = [('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#8B0000')),
                 ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#D4A017')),
                 ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                 ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]
        for index, record in enumerate(rows, 1):
            background = '#EEEEEE' if record.get('Status') == 'Absent' else '#FFF3CD' if record.get('Status') == 'Incomplete' else '#FFFFFF'
            style.append(('BACKGROUND', (0, index), (-1, index), colors.HexColor(background)))
        table.setStyle(TableStyle(style))
        story.append(table)
    document.build(story)
    return output.getvalue()


def export_lunch_xlsx(title, subtitle, tables, logo_path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Lunch report'
    count = len(tables[0][1]) if tables else len(DAILY_COLUMNS)
    for index, text in enumerate([title, subtitle, SCHEDULE,
                                  'Blank clockings mean no scan was recorded. Net hours require all four clockings.'], 1):
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=count)
        cell = sheet.cell(index, 1, text)
        cell.font = Font(name='Arial', bold=index == 1, size=14 if index == 1 else 10, color='8B0000')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.row_dimensions[index].height = 25
    border = Border(*(Side(style='thin', color='D4A017') for _ in range(4)))
    current = 6
    for heading, columns, rows in tables:
        sheet.merge_cells(start_row=current, start_column=1, end_row=current, end_column=len(columns))
        cell = sheet.cell(current, 1, heading)
        cell.font = Font(name='Arial', bold=True, color='8B0000')
        current += 1
        for index, label in enumerate(columns, 1):
            cell = sheet.cell(current, index, label)
            cell.font = Font(name='Arial', bold=True, color='FFF3CD')
            cell.fill = PatternFill('solid', fgColor='8B0000')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[current].height = 30
        current += 1
        for record in rows:
            for index, label in enumerate(columns, 1):
                value = record.get(label, '')
                # Force names/IDs to text; terminal fields must never become formulas.
                cell = sheet.cell(current, index, value)
                if isinstance(value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(wrap_text=True, vertical='center')
                cell.border = border
                fill = 'EEEEEE' if record.get('Status') == 'Absent' else 'FFF3CD' if record.get('Status') == 'Incomplete' else 'FFFFFF'
                cell.fill = PatternFill('solid', fgColor=fill)
            sheet.row_dimensions[current].height = 30
            current += 1
        for index, label in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 32 if label == 'Employee Name' else 35 if label == 'Missing Clockings' else 14
        current += 1
    sheet.freeze_panes = 'D8'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_lunch_report(events, site, start_date, end_date, tag, logo_path=None):
    import pandas as pd
    import streamlit as st
    rows = parse_lunch_events(events, start_date, end_date)
    daily, summary = lunch_tables(rows, site, start_date, end_date)
    st.markdown('### Buguruni report with lunch time')
    st.caption(SCHEDULE)
    st.info('Clock four times: arrival at 8 am, lunch out at 1 pm, lunch back at 2 pm, and departure at 4 pm. '
            'Missing clockings stay blank. Net hours are calculated only when all four clockings are present.')
    days = (end_date - start_date).days + 1
    m1, m2, m3, m4 = st.columns(4)
    m1.metric('Days', days)
    m2.metric('Staff Roster', sum(len(table[2]) for table in summary))
    m3.metric('Complete Days', sum(row['complete'] for row in rows))
    m4.metric('Incomplete Days', sum(not row['complete'] for row in rows))
    downloads, viewer = st.tabs(['Download Reports', 'View Reports'])
    with viewer:
        daily_view, summary_view = st.tabs(['Daily Attendance with Lunch', 'Period Summary with Lunch'])
        for target, tables in [(daily_view, daily), (summary_view, summary)]:
            with target:
                for heading, columns, records in tables:
                    st.markdown(f'**{heading}**')
                    st.dataframe(pd.DataFrame(records, columns=columns), hide_index=True, use_container_width=True)
    subtitle = f"Site: Buguruni | {start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')}"
    with downloads:
        daily_tab, summary_tab = st.tabs(['Daily with Lunch', 'Summary with Lunch'])
        for target, label, tables in [(daily_tab, 'Daily', daily), (summary_tab, 'Summary', summary)]:
            with target:
                columns = st.columns(3)
                exporters = [('Word (.docx)', 'docx', export_lunch_docx,
                              'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
                             ('PDF', 'pdf', export_lunch_pdf, 'application/pdf'),
                             ('Excel (.xlsx)', 'xlsx', export_lunch_xlsx,
                              'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')]
                for column, (name, extension, exporter, mime) in zip(columns, exporters):
                    data = exporter(f'BUGURUNI {label.upper()} REPORT WITH LUNCH TIME', subtitle, tables, logo_path)
                    with column:
                        st.download_button(f'Download {name}', data=data,
                                           file_name=f'Buguruni_{label}_With_Lunch_{tag}.{extension}', mime=mime)
